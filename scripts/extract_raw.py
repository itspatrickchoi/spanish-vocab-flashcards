"""
Step 1: Raw, lossless extraction of every data row from the source workbook.

Reads with openpyxl's default settings, which does NOT respect Excel's
AutoFilter hidden-row state -- every row, hidden or not, is captured. This
is the fix for the "AI skims and misses words" problem: the source file has
434 rows hidden by a saved filter, 15 of which contain real content.

Output: data/extracted_rows.json, one record per data row (row 2..N), with
row number, raw A/B/C values, and whether Excel had it marked hidden.
"""

import json
from pathlib import Path

import openpyxl

SOURCE = Path(__file__).parent.parent / "data" / "source" / "spanish_vocab_source.xlsx"
OUT = Path(__file__).parent.parent / "data" / "extracted_rows.json"


def cell_text(value):
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def main():
    wb = openpyxl.load_workbook(SOURCE, data_only=True)
    ws = wb["Sheet1"]

    records = []
    max_row = ws.max_row
    for row_idx in range(2, max_row + 1):  # skip header row 1
        row_dim = ws.row_dimensions.get(row_idx)
        hidden = bool(row_dim.hidden) if row_dim is not None else False

        a = cell_text(ws.cell(row=row_idx, column=1).value)
        b = cell_text(ws.cell(row=row_idx, column=2).value)
        c = cell_text(ws.cell(row=row_idx, column=3).value)

        is_empty = a is None and b is None and c is None

        records.append(
            {
                "row": row_idx,
                "front_raw": a,
                "back_raw": b,
                "overflow_raw": c,
                "hidden": hidden,
                "is_empty": is_empty,
            }
        )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")

    total = len(records)
    empty = sum(1 for r in records if r["is_empty"])
    non_empty = total - empty
    hidden_nonempty = sum(1 for r in records if r["hidden"] and not r["is_empty"])

    print(f"Total data rows: {total} (rows 2..{max_row})")
    print(f"Empty: {empty}")
    print(f"Non-empty: {non_empty}")
    print(f"Hidden rows with real content: {hidden_nonempty}")
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
