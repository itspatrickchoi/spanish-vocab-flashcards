"""
Step 3: Build the human review workbook.

Three sheets:
  - Summary: row-accounting reconciliation + top-line counts.
  - Review: every drafted card (Front/Back/Notes/Reversible/Confidence) next
    to its source row(s) and original text, for the user to check/edit.
  - Excluded: every row dropped as noise/duplicate/blank, with a reason.

No formulas are needed here (this is a data review sheet, not a financial
model), so no recalc pass is required -- but formatting follows the same
"professional, readable" bar: Arial, frozen header, wrapped long text
columns, and light color flags for confidence so low/medium items are easy
to spot while skimming ~500 rows.
"""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent.parent
DRAFTS = json.load(open(ROOT / "data" / "merged_drafts.json", encoding="utf-8"))
CLASSIFIED = json.load(open(ROOT / "data" / "classified_rows.json", encoding="utf-8"))
OUT = ROOT / "data" / "review_spanish_flashcards.xlsx"

FONT_NAME = "Arial"
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
BODY_FONT = Font(name=FONT_NAME)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")

CONF_FILL = {
    "low": PatternFill("solid", fgColor="F8CBAD"),
    "medium": PatternFill("solid", fgColor="FFE699"),
    "high": None,
}


def rows_str(rows):
    return ", ".join(str(r) for r in rows)


def style_header(ws, headers):
    ws.freeze_panes = "A2"
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main():
    wb = Workbook()

    # ---------------- Summary sheet ----------------
    ws = wb.active
    ws.title = "Summary"
    counts = CLASSIFIED["counts"]
    excl_cats = {}
    for e in CLASSIFIED["excluded_rows"]:
        excl_cats[e["category"]] = excl_cats.get(e["category"], 0) + len(e["rows"])

    flagged = [d for d in DRAFTS if d.get("flag_reason")]
    reversible = [d for d in DRAFTS if d.get("reversible")]
    conf_counts = {"high": 0, "medium": 0, "low": 0}
    for d in DRAFTS:
        conf_counts[d.get("confidence", "high")] = conf_counts.get(d.get("confidence", "high"), 0) + 1

    lines = [
        ("Source rows (data rows, excl. header)", counts["total_data_rows"]),
        ("Rows accounted for", counts["physical_rows_accounted_for"]),
        ("Reconciles to source row count", "YES" if counts["reconciles"] else "NO -- CHECK classify_rows.py"),
        ("", ""),
        ("Draft flashcards produced", len(DRAFTS)),
        ("  - marked reversible (also quizzed meaning->word)", len(reversible)),
        ("  - recognition-only (word->meaning only)", len(DRAFTS) - len(reversible)),
        ("  - confidence: high", conf_counts.get("high", 0)),
        ("  - confidence: medium", conf_counts.get("medium", 0)),
        ("  - confidence: low", conf_counts.get("low", 0)),
        ("  - flagged for your review (see Review sheet)", len(flagged)),
        ("", ""),
        ("Rows excluded (noise/duplicate/blank)", sum(excl_cats.values())),
    ]
    for cat, n in sorted(excl_cats.items(), key=lambda x: -x[1]):
        lines.append((f"  - {cat}", n))

    ws.cell(row=1, column=1, value="Spanish Vocab -> Anki: build summary").font = Font(name=FONT_NAME, bold=True, size=14)
    r = 3
    for label, val in lines:
        ws.cell(row=r, column=1, value=label).font = Font(name=FONT_NAME, bold=label.strip() != "" and not label.startswith("  "))
        ws.cell(row=r, column=2, value=val).font = BODY_FONT
        r += 1
    set_widths(ws, [48, 14])

    # ---------------- Review sheet ----------------
    ws = wb.create_sheet("Review")
    headers = [
        "Source row(s)", "Category", "Front (Spanish)", "Back (Spanish def. + example)",
        "Notes (EN gloss + usage)", "Reversible?", "Confidence", "Flag - please check", "Original source text",
    ]
    style_header(ws, headers)
    set_widths(ws, [12, 16, 22, 42, 34, 11, 11, 34, 42])

    drafts_sorted = sorted(DRAFTS, key=lambda d: min(d["rows"]))
    for i, d in enumerate(drafts_sorted, start=2):
        vals = [
            rows_str(d["rows"]),
            d["category"],
            d["front"],
            d["back"],
            d["notes"],
            "Yes" if d.get("reversible") else "No",
            d.get("confidence", "high"),
            d.get("flag_reason") or "",
            d["source_text"],
        ]
        for col, v in enumerate(vals, start=1):
            c = ws.cell(row=i, column=col, value=v)
            c.font = BODY_FONT
            c.alignment = WRAP if col in (4, 5, 8, 9) else TOP
        fill = CONF_FILL.get(d.get("confidence", "high"))
        if fill:
            for col in range(1, len(headers) + 1):
                ws.cell(row=i, column=col).fill = fill
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(drafts_sorted) + 1}"

    # ---------------- Excluded sheet ----------------
    ws = wb.create_sheet("Excluded")
    headers = ["Source row(s)", "Category", "Reason", "Original text (row 1 of group)"]
    style_header(ws, headers)
    set_widths(ws, [14, 20, 46, 50])
    excl_sorted = sorted(CLASSIFIED["excluded_rows"], key=lambda e: min(e["rows"]))
    for i, e in enumerate(excl_sorted, start=2):
        raw0 = e["raw"][0]
        raw_preview = raw0.get("front_raw") or raw0.get("back_raw") or raw0.get("overflow_raw") or ""
        vals = [rows_str(e["rows"]), e["category"], e["reason"], raw_preview]
        for col, v in enumerate(vals, start=1):
            c = ws.cell(row=i, column=col, value=v)
            c.font = BODY_FONT
            c.alignment = WRAP if col == 4 else TOP
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(excl_sorted) + 1}"

    wb.save(OUT)
    print(f"Wrote {OUT}")
    print(f"Review rows: {len(drafts_sorted)}, Excluded rows: {sum(len(e['rows']) for e in excl_sorted)}")


if __name__ == "__main__":
    main()
